"""Build the illustrated Word manual from its reviewable Markdown source.

Run with the Codex bundled Python runtime (python-docx and Pillow required).
The source and original browser JPEGs remain the review authority.
"""
import argparse
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
from lxml import etree


CHINESE_FONT = 'Microsoft YaHei'


def configure_chinese_fonts(document):
    """Keep theme fonts from overriding Chinese text in Word and PDF exports."""
    namespaces = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                  'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    japanese_fonts = {'MS Gothic', 'MS-Gothic', 'ＭＳ ゴシック', 'MS Mincho', 'ＭＳ 明朝'}
    for part in document.part.package.parts:
        name = str(part.partname)
        if not name.startswith('/word/') or not name.endswith('.xml'):
            continue
        element = getattr(part, '_element', None)
        root = element if element is not None else etree.fromstring(part.blob)
        # Include text in table cells, headers, footers and field results.
        for run in root.findall('.//w:r', namespaces):
            properties = run.find(qn('w:rPr'))
            if properties is None:
                properties = OxmlElement('w:rPr')
                run.insert(0, properties)
            if properties.find(qn('w:rFonts')) is None:
                properties.insert(0, OxmlElement('w:rFonts'))
        for fonts in root.findall('.//w:rFonts', namespaces):
            latin = 'Consolas' if 'Consolas' in fonts.attrib.values() else CHINESE_FONT
            for key in list(fonts.attrib):
                if 'theme' in key.lower():
                    del fonts.attrib[key]
            for script in ['ascii', 'hAnsi', 'cs']:
                fonts.set(qn('w:' + script), latin)
            fonts.set(qn('w:eastAsia'), CHINESE_FONT)
        for tag in ['w:lang', 'w:themeFontLang']:
            for language in root.findall('.//' + tag, namespaces):
                language.set(qn('w:eastAsia'), 'zh-CN')
        for tag in ['a:latin', 'a:ea', 'a:cs']:
            for font in root.findall('.//' + tag, namespaces):
                font.set('typeface', CHINESE_FONT)
        for font in root.findall('.//a:font', namespaces):
            if font.get('script') in {'Jpan', 'Hans', 'Hant'} or font.get('typeface') in japanese_fonts:
                font.set('typeface', CHINESE_FONT)
        for font in list(root.findall('.//w:font', namespaces)):
            if font.get(qn('w:name')) in japanese_fonts:
                font.getparent().remove(font)
        if name == '/word/fontTable.xml' and not any(
                font.get(qn('w:name')) == CHINESE_FONT for font in root):
            font = OxmlElement('w:font')
            font.set(qn('w:name'), CHINESE_FONT)
            alias = OxmlElement('w:altName')
            alias.set(qn('w:val'), '微软雅黑')
            font.append(alias)
            root.append(font)
        if element is None:
            part._blob = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)


def inline(paragraph, text):
    """Keep emphasis while displaying link labels and addresses as readable text."""
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1（\2）', text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(part.strip('*`') if part.startswith(('**', '`')) else part)
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Consolas'


def build(source, destination):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.64), Inches(.62)
    section.left_margin, section.right_margin = Inches(.7), Inches(.7)
    section.header_distance, section.footer_distance = Inches(.25), Inches(.25)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        style = doc.styles[name]
        style.font.name = CHINESE_FONT
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CHINESE_FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
        grid = OxmlElement('w:snapToGrid')
        grid.set(qn('w:val'), '0')
        style.element.get_or_add_pPr().append(grid)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(27)
    doc.styles['Title'].paragraph_format.space_after = Pt(14)
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].paragraph_format.space_after = Pt(8)
    doc.styles['Caption'].paragraph_format.line_spacing = 1.05
    doc.styles['Caption'].paragraph_format.keep_with_next = False
    header = section.header.paragraphs[0]
    header.text = '专题数字图书馆  使用手册'
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    version = re.search(r'文档版本为 ([0-9.]+)', source.read_text(encoding='utf-8'))
    footer.add_run('版本 ' + (version.group(1) if version else '1.0') + '  ·  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.title = '专题数字图书馆使用手册'
    doc.core_properties.subject = '浏览 搜索 上传 元数据核对 证据问答 操作说明'
    doc.core_properties.author = '数字图书馆项目'
    doc.core_properties.last_modified_by = '数字图书馆项目'
    lines = source.read_text(encoding='utf-8').splitlines()
    index, chapter, code, next_page = 0, 1, False, False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            next_page = True
            continue
        if line.startswith('```'):
            code = not code
            continue
        if code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            run = p.add_run(line)
            run.font.name, run.font.size = 'Consolas', Pt(10)
            continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
            continue
        if line.startswith('## '):
            chapter = int(line.split()[1])
            p = doc.add_paragraph(line[3:], 'Heading 1')
            p.paragraph_format.page_break_before = next_page
            next_page = False
            continue
        image_match = re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)', line)
        if image_match:
            caption, rel = image_match.groups()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            filename = Path(rel).name
            narrow = {'01b-classification.jpg', '02-covers.jpg', '03-upload.jpg',
                      '04-recognizing.jpg', '13-citation.jpg', '14-source-page.jpg'}
            medium = {'06a-metadata-year.jpg', '15b-scanned-body.jpg', '16-system-status.jpg'}
            width = 5.4 if filename in narrow else (6.2 if filename in medium else 7.0)
            pic = p.add_run().add_picture(str(source.parent / rel), width=Inches(width))
            pic._inline.docPr.set('descr', caption)
            # Crop only in Word layout; the source browser JPEG remains unchanged.
            crop = ({'l': 28000, 'r': 6000, 't': 55000, 'b': 0}
                    if rel.endswith('04-recognizing.jpg') else
                    {'l': 28000, 'r': 6000, 't': 20000, 'b': 25000}
                    if rel.endswith('15b-scanned-body.jpg') else None)
            if crop:
                rect = OxmlElement('a:srcRect')
                for key, value in crop.items():
                    rect.set(key, str(value))
                pic._inline.graphic.graphicData.pic.blipFill.insert(1, rect)
                with Image.open(source.parent / rel) as original:
                    img_w, img_h = original.size
                visible_w = img_w * (1 - (crop['l'] + crop['r']) / 100000)
                visible_h = img_h * (1 - (crop['t'] + crop['b']) / 100000)
                pic.height = Inches(width * visible_h / visible_w)
                caption += ' 局部'
            p = doc.add_paragraph(caption, 'Caption')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip())
                index += 1
            cells = [[x.strip() for x in row.strip('|').split('|')] for row in rows]
            cells = [r for r in cells if not all(re.fullmatch(r':?-+:?', x) for x in r)]
            table = doc.add_table(rows=0, cols=len(cells[0]))
            table.autofit = False
            widths = ([1.15, 3.05, 2.9] if len(cells[0]) == 3 else [1.7, 5.4])
            for col, width in zip(table.columns, widths):
                col.width = Inches(width)
            for ri, values in enumerate(cells):
                row = table.add_row()
                if ri == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                cant_split = OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(cant_split)
                for ci, (cell, value) in enumerate(zip(row.cells, values)):
                    cell.width = Inches(widths[ci])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for edge in ['top', 'left', 'bottom', 'right']:
                        edge_el = OxmlElement('w:' + edge)
                        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                            edge_el.set(qn('w:' + key), val)
                        borders.append(edge_el)
                    tcpr.append(borders)
                    margins = OxmlElement('w:tcMar')
                    for edge, size in [('top', '85'), ('bottom', '85'), ('left', '105'), ('right', '105')]:
                        el = OxmlElement('w:' + edge)
                        el.set(qn('w:w'), size)
                        el.set(qn('w:type'), 'dxa')
                        margins.append(el)
                    tcpr.append(margins)
                    fill = OxmlElement('w:shd')
                    fill.set(qn('w:fill'), 'E8EEF3' if ri == 0 else ('FFFFFF' if ri % 2 else 'F7F9FB'))
                    tcpr.append(fill)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.06
                    inline(p, value)
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.bold = ri == 0
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        p = doc.add_paragraph()
        if re.match(r'^\d+\. ', line):
            p.paragraph_format.left_indent = Inches(.18)
            p.paragraph_format.first_line_indent = Inches(-.18)
        inline(p, line)
    destination.parent.mkdir(parents=True, exist_ok=True)
    configure_chinese_fonts(doc)
    doc.save(destination)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())
